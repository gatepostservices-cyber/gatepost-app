#!/usr/bin/env python3
"""Index Gatepost business files into public.doc_chunks (Ask AI search).
Walks a local mirror of OneDrive Gatepost/ (default /workspace/gatepost-mirror; refresh it first
with the OneDrive connector), extracts text from docx/xlsx/pdf/txt/md/csv, chunks it, and upserts
through the token-gated RPC post_doc_file (token: ~/.config/gatepost/request-token). No service key.
  python3 tools/index-files.py [MIRROR_ROOT] [--no-prune]
"""
import os, sys, re, json, csv, subprocess, datetime, urllib.request
ROOT = next((a for a in sys.argv[1:] if not a.startswith('--')), '/workspace/gatepost-mirror')
DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
cfg = open(os.path.join(DIR, 'config.js')).read()
URL = re.search(r'SUPABASE_URL:\s*"([^"]+)"', cfg).group(1); KEY = re.search(r'SUPABASE_ANON_KEY:\s*"([^"]+)"', cfg).group(1)
TOKEN = open(os.environ.get('GATEPOST_REQUEST_TOKEN_FILE', os.path.expanduser('~/.config/gatepost/request-token'))).read().strip()
SKIP_DIRS = ('99_Archive/Duplicates', '04_Branding', '08_Photos')

def rpc(fn, body):
    req = urllib.request.Request(f'{URL}/rest/v1/rpc/{fn}', json.dumps(body).encode(), {'apikey': KEY, 'Content-Type': 'application/json'})
    return json.load(urllib.request.urlopen(req, timeout=60))

def extract(p):
    ext = p.rsplit('.', 1)[-1].lower()
    if ext == 'docx':
        import docx; d = docx.Document(p); out = [x.text for x in d.paragraphs if x.text.strip()]
        for t in d.tables:
            for r in t.rows:
                cells = []
                for c in r.cells:
                    if c.text.strip() and c.text.strip() not in cells: cells.append(c.text.strip())
                if cells: out.append(' | '.join(cells))
        return '\n'.join(out)
    if ext == 'xlsx':
        import openpyxl; wb = openpyxl.load_workbook(p, data_only=True, read_only=True); out = []
        for ws in wb.worksheets:
            out.append(f'## Sheet: {ws.title}')
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v not in (None, '')]
                if vals: out.append(' | '.join(vals))
        return '\n'.join(out)
    if ext == 'pdf': return subprocess.run(['pdftotext', '-layout', p, '-'], capture_output=True, text=True).stdout
    if ext == 'csv': return '\n'.join(' | '.join(r) for r in csv.reader(open(p, errors='replace')))
    if ext in ('txt', 'md'): return open(p, errors='replace').read()
    return None

def chunks(text, size=1500, overlap=200):
    text = re.sub(r'[ \t]+', ' ', re.sub(r'\n{3,}', '\n\n', text)).strip(); out = []; i = 0
    while i < len(text):
        j = min(len(text), i + size)
        if j < len(text):
            k = text.rfind('\n', i + size // 2, j); j = k if k > 0 else j
        out.append(text[i:j].strip()); i = j if j >= len(text) else max(j - overlap, i + 1)
    return [c for c in out if c]

files = chunk_total = 0; skipped = []; keep = []
for dp, dn, fn in os.walk(ROOT):
    dn.sort()
    for f in sorted(fn):
        full = os.path.join(dp, f); rel = os.path.relpath(full, ROOT)
        if any(s in rel for s in SKIP_DIRS): skipped.append((rel, 'skipped folder')); continue
        try: text = extract(full)
        except Exception as e: skipped.append((rel, f'error: {e}')); continue
        if text is None: skipped.append((rel, 'unsupported type')); continue
        if len(text.strip()) < 20: skipped.append((rel, 'no text')); continue
        cs = chunks(text); mod = datetime.datetime.fromtimestamp(os.path.getmtime(full), datetime.timezone.utc).isoformat()
        title = os.path.splitext(f)[0].replace('_', ' ')
        n = rpc('post_doc_file', {'p_token': TOKEN, 'p_path': rel, 'p_title': title, 'p_modified': mod, 'p_chunks': cs})
        files += 1; chunk_total += n; keep.append(rel)
if '--no-prune' not in sys.argv and keep: print('pruned', rpc('prune_doc_files', {'p_token': TOKEN, 'p_keep': keep}), 'stale chunks')
print(f'indexed {files} files, {chunk_total} chunks')
for r, why in skipped: print('  skipped:', r, '-', why)
